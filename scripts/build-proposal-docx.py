from pathlib import Path
from io import BytesIO
from docx import Document
from PIL import Image
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'site/designs/downloads/onlyyuuka-website-proposal.docx'
doc=Document(); sec=doc.sections[0]
sec.page_width=Mm(210);sec.page_height=Mm(297);sec.top_margin=Mm(19);sec.bottom_margin=Mm(18);sec.left_margin=Mm(23);sec.right_margin=Mm(23)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
 s=doc.styles[name];s.font.name='Noto Sans CJK JP';s.font.color.rgb=RGBColor(0,0,0);s.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Noto Sans CJK JP')
 s.paragraph_format.space_after=Pt(6)
 snap=OxmlElement('w:snapToGrid');snap.set(qn('w:val'),'0');s.element.get_or_add_pPr().append(snap)
normal=doc.styles['Normal'];normal.font.size=Pt(10);normal.paragraph_format.line_spacing=Pt(15)
for name,size in [('Title',27),('Subtitle',12),('Heading 1',21),('Heading 2',13),('Heading 3',11)]:
 doc.styles[name].font.size=Pt(size);doc.styles[name].font.bold=name.startswith('Heading');doc.styles[name].paragraph_format.space_before=Pt(8);doc.styles[name].paragraph_format.keep_with_next=True
for name in ['Title','Heading 1']:
 doc.styles[name].font.name='Noto Serif CJK JP';doc.styles[name].element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Noto Serif CJK JP')
doc.styles['Title'].paragraph_format.line_spacing=Pt(35)
doc.styles['Subtitle'].font.italic=False
doc.styles['Heading 1'].paragraph_format.line_spacing=Pt(29)
doc.styles['Heading 2'].paragraph_format.line_spacing=Pt(20)
doc.styles['Caption'].font.size=Pt(8);doc.styles['Caption'].font.color.rgb=RGBColor.from_string('666666');doc.styles['Caption'].paragraph_format.line_spacing=1.25
props=doc.core_properties;props.title='onlyyuuka ホームページ制作提案書';props.subject='3つのデザイン案と掲載内容および準備事項';props.author='';props.last_modified_by='';props.keywords='onlyyuuka ホームページ デザイン 提案'
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=footer.add_run('onlyyuuka  |  ');r.font.size=Pt(8);r.font.color.rgb=RGBColor.from_string('777777')
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');footer._p.append(f)
def p(t,style=None):return doc.add_paragraph(t,style)
pending_page=False
def h(t):
 global pending_page
 para=doc.add_heading(t,1);para.paragraph_format.page_break_before=pending_page;pending_page=False;return para
def sub(t):return doc.add_heading(t,2)
def page():
 global pending_page
 pending_page=True
def link(text,url):
 para=doc.add_paragraph(); hyp=OxmlElement('w:hyperlink');hyp.set(qn('r:id'),para.part.relate_to(url,RT.HYPERLINK,is_external=True));r=OxmlElement('w:r');pr=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'4F6654');pr.append(color);r.append(pr);t=OxmlElement('w:t');t.text=text;r.append(t);hyp.append(r);para._p.append(hyp);return para

def table(headers,rows,widths):
 t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 for col,w in zip(t.columns,widths):col.width=Mm(w)
 for c,w,title in zip(t.rows[0].cells,widths,headers):c.width=Mm(w);c.text=title
 trPr=t.rows[0]._tr.get_or_add_trPr();repeat=OxmlElement('w:tblHeader');trPr.append(repeat)
 for row in rows:
  cells=t.add_row().cells
  for c,w,txt in zip(cells,widths,row):c.width=Mm(w);c.text=txt
 for i,row in enumerate(t.rows):
  trpr=row._tr.get_or_add_trPr();nosplit=OxmlElement('w:cantSplit');trpr.append(nosplit)
  for j,c in enumerate(row.cells):
   c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;tcpr=c._tc.get_or_add_tcPr()
   borders=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right','insideH','insideV']:
    e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
   tcpr.append(borders);m=OxmlElement('w:tcMar')
   for edge,twips in [('top',80),('bottom',80),('left',110),('right',110)]:
    e=OxmlElement('w:'+edge);e.set(qn('w:w'),str(twips));e.set(qn('w:type'),'dxa');m.append(e)
   tcpr.append(m)
   sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'404040' if i==0 else ('F6F6F6' if i%2==0 else 'FFFFFF'));tcpr.append(sh)
   for para in c.paragraphs:
    para.paragraph_format.line_spacing=Pt(13);para.paragraph_format.space_after=Pt(0)
    for r in para.runs:r.font.size=Pt(9);r.font.bold=i==0;r.font.color.rgb=RGBColor(255,255,255) if i==0 else RGBColor(0,0,0)
 return t

def picture(file,caption):
 para=doc.add_paragraph();para.paragraph_format.space_before=Pt(5);para.paragraph_format.space_after=Pt(5)
 para.paragraph_format.line_spacing=1.0
 source=ROOT/'site/designs/assets'/file
 if not source.exists(): source=source.with_suffix('.webp')
 img=Image.open(source);iw,ih=img.size;para.alignment=WD_ALIGN_PARAGRAPH.CENTER
 buffer=BytesIO();img.save(buffer,format='PNG');buffer.seek(0)
 shape=para.add_run().add_picture(buffer,width=Mm(min(164,80*iw/ih))); shape._inline.docPr.set('descr',caption);para.paragraph_format.keep_with_next=True;p(caption,'Caption')

p('onlyyuuka ご担当者様')
p('onlyyuuka\nホームページ制作提案書','Title')
p('3つのデザイン案と掲載内容および準備事項','Subtitle')
p('2026年9月7日')
sub('今回のご提案')
p('作品の魅力と作り手の考えが伝わり、問い合わせにつながるホームページをご提案します。商品を直接販売するECサイトではなく、活動や作品を紹介する名刺サイトとして設計します。')
p('基本構成はホーム、作り手紹介、作品、お問い合わせの4ページです。お便りの登録欄を組み合わせ、初めて知った方とも継続的につながれる形を想定します。')
sub('まずご覧いただきたいこと')
p('3案は、同じ情報を色違いにしたものではありません。作品、作り手、素材の変化のうち、何を最初に伝えるかを変えています。今回は、作品を中心に見せやすい案01を基本案としておすすめします。')
table(['案','主役','伝えたい印象'],[['01 ギャラリー','作品の質感と存在感','上品に作品を鑑賞する'],['02 アトリエ','手仕事と作り手の視点','制作する人に親しむ'],['03 編集誌','素材から作品への変化','リメイクの面白さを知る']],[37,58,69])
link('3つのデザイン案を見る','https://yuearmy.github.io/mock-up/designs/')
p('本書でご相談したいのは、デザインの方向性と、最初に掲載する作品・内容です。正式な制作範囲、費用、日程は、ご希望と運用方法を伺ってから決めます。')

for num,title,lead,file,caption,order,focus,materials,slug in [
 ('01','布の小さなギャラリー','作品を一枚ずつ眺めるように、写真と余白で魅力を伝える案です。','gallery-textile.png','方向性のイメージ写真　実作品の記録ではありません','作品の世界観 → 代表作品 → 作り手 → お問い合わせ','作品写真を大きく配置し、文字を絞ります。作品ごとの素材や使う場面へ自然に進めるため、展示・取材・協業を検討する方にも活動を伝えやすい構成です。','高解像度の代表作品、布の質感、展示や使用場面の写真。写真の完成度が印象を大きく左右します。','gallery'),
 ('02','作り手のアトリエ','制作している人の視点を入口に、手仕事への親しみを育てる案です。','hands-process.png','制作風景のイメージ写真　正式掲載時はご本人の手元写真に差し替え','制作の手元 → 作り手の視点 → 工程 → 作品 → 相談','手元の写真にノートのような文章を重ね、素材を選ぶことや仕立てる過程を紹介します。作品だけでなく、誰がどんな気持ちでつくるかを伝えたい場合に適しています。','ご本人の手元や制作場所、制作中の写真、始めたきっかけや大切にしていることを話した短いメモ。','atelier'),
 ('03','布をめぐる編集誌','素材の名前から作品をたどり、リメイクの変化を読むように楽しむ案です。','still-life.png','暮らしの中の布小物のイメージ写真　実作品の使用写真に差し替える想定','素材の変化 → 作品ごとの短い記事 → 作り手 → 相談','大きな誌名、罫線、記事番号を使い、帯からテーブルランナー、黒帯からバッグといった変化を説明します。見た目と使い道の両方を知ってもらいたい場合に適しています。','同じ素材の制作前後が分かる写真、使っている場面、どこを残し何を変えたかの短い解説。','journal')]:
 page();h(f'案{num} {title}');p(lead);picture(file,caption);sub('見せ方と向いている目的');p(focus);sub('ページを読む順番');p(order);sub('特に用意したい素材');p(materials);link('このデザイン案を開く',f'https://yuearmy.github.io/mock-up/designs/{slug}/')

page();h('掲載するコンテンツ');p('初回は次の内容を中心に構成します。下記の字数や枚数は目安で、支給いただく素材に合わせて調整します。')
table(['掲載場所','載せる内容','必要な情報の粒度'],[
 ['ホーム','ブランドと第一印象','正式な名前、肩書き、短い見出し、活動説明2〜3行、代表写真、作品と相談への入口'],
 ['ホーム','制作の考え方','残したい素材の特徴や、リメイクへの考えを100〜150字程度'],
 ['作品','代表作品3〜6点','作品名、写真、素材、寸法、制作年、使う場面、工夫した点を1〜2文'],
 ['作品詳細','写真と短い制作解説','全体、裏や内側、質感、使用場面。素材から何に変わったか、問い合わせへの入口'],
 ['作り手紹介','活動と制作背景','何をつくる人か、始めたきっかけ、大切にしていること。200〜400字程度と本人または手元写真'],
 ['お問い合わせ','受け付ける相談と入力欄','作品、オーダー、取材、展示、コラボなど。名前、メール、相談種別、相談内容'],
 ['ページ下部','ニュースレター','届く内容、配信頻度の目安、登録と解除の案内。配信の運用が決まってから設置'],
 ['フッター','公式リンクと運営案内','公式SNS、連絡窓口、運営者表記、個人情報の取り扱いへの入口']],[27,45,92])
sub('素材がそろったら追加できる内容')
p('同じ作品のリメイク前後、展示やイベントのお知らせ、確認済みの掲載実績、持ち込みオーダーの詳しい案内などを追加できます。更新が必要な内容は、担当者を決めたうえで掲載します。')

page();h('準備していただきたいもの');p('最初は完成した原稿でなくても大丈夫です。写真と短いメモをもとに、サイトに合う文章へ整えます。')
table(['準備するもの','目安','ほしい内容'],[
 ['代表作品の情報','3〜6作品','作品名、素材、寸法、制作年、使う場面、残した特徴や仕立ての工夫'],
 ['作品写真','1作品3〜5枚','全体、裏や内側、布の質感、サイズが分かる写真、使っている場面'],
 ['メイン写真','横と縦 各1枚','代表作品を自然光で撮影。周囲に余白があり、長辺2,000px以上の元画像が目安'],
 ['作り手と工程の写真','本人等1〜2枚\n工程3〜5枚','本人の写真または実際の手元、元の布、素材選び、裁断、縫製、仕上げ'],
 ['作家紹介のメモ','短い箇条書きで可','始めたきっかけ、好きな素材、大切にしていること、これから取り組みたいこと'],
 ['連絡先と公式SNS','窓口を確定','公開する窓口、フォーム受信先、対応する人、公式InstagramのURL']],[40,34,90])
sub('今回のサンプルと正式掲載の違い')
p('作品一覧には既存資料の作品写真を使用しています。大きく見せるための高解像度写真と、寸法・制作年などの情報をそろえてから正式掲載します。')
p('布の展示、手元、机上の写真はデザインの方向を伝えるイメージです。実作品や本人の制作風景を示す写真とは区別し、正式公開前に掲載対象と利用許可を確認します。')
sub('確認の分担')
p('作品や活動の事実、写真の使用許可はクライアント側でご確認いただきます。制作者は、原稿の整理、写真の配置、スマホでの読みやすさ、リンクやフォームの動作を整えます。')

page();h('公開までに決めること');p('画面をつくることと、実際に問い合わせを受け付けて更新することは、それぞれ準備が必要です。次の項目を打ち合わせで確認します。')
table(['項目','決めたい内容'],[
 ['制作範囲と予算','初回のページと内容、追加機能、希望公開時期。制作費と維持費の対象を確認'],
 ['公開先と管理','使うURL、ドメイン、契約名義、管理アカウントの所有者、更新する人'],
 ['お問い合わせ','受け付ける相談、窓口と受信先、対応担当、返信目安、迷惑送信への対応'],
 ['ニュースレター','実施の有無、内容、配信する人と頻度、サービス、同意・解除・登録者管理'],
 ['掲載情報と確認','公開する名前、写真、原稿、実績、個人情報の扱い、最終承認の担当']],[39,125])
sub('進め方')
p('まずデザインの方向と代表作品を選びます。次に写真と原稿をそろえ、ページへ反映します。受付方法と運用が決まったら本番のフォームなどを接続し、PCとスマホ、実際の送受信を確認してから公開します。')
sub('今回のサンプルで確認できること')
p('3案の各4ページ、作品カテゴリ、作品詳細、作品名を引き継ぐ相談、入力内容の確認と修正を試せます。Computer UseでPC・スマホ・中間サイズの表示と主要な操作を確認しています。問い合わせの送信、メルマガ登録、購入や決済は行いません。')
sub('最初の打ち合わせで')
p('代表作品を3つ選び、「どんな人に、何を伝えたいか」をお聞かせください。それをもとに、案01・02・03のどれを軸にするか、必要な写真と内容を一緒に決めていきます。')
link('デザイン3案と準備リストを見る','https://yuearmy.github.io/mock-up/designs/')
# Remove inherited theme font overrides and decorative title borders.
for elem in doc.styles.element.iter():
 if elem.tag==qn('w:rFonts'):
  for key in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
   elem.attrib.pop(qn('w:'+key),None)
for elem in list(doc.styles.element.iter(qn('w:pBdr'))): elem.getparent().remove(elem)
for elem in list(doc.element.iter(qn('w:pBdr'))): elem.getparent().remove(elem)
OUT.parent.mkdir(parents=True,exist_ok=True);doc.save(OUT);print(OUT)
